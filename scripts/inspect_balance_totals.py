#!/usr/bin/env python3
"""Inspect specific balance accounts (512xxx, 46xxx, 58xxx) and totals row."""
from openpyxl import load_workbook

VERIF = "/home/z/my-project/upload/Verif BG 30092026 V2.xlsx"
wb = load_workbook(VERIF, data_only=True, read_only=True)
ws = wb["Balance des comptes"]

# Columns:
# 0: N° compte (merged across 1-3)
# 4: Intitulé (merged across 4-8)
# 9: Mouvements Débit (merged 9-10)
# 11: Mouvements Crédit (merged 11-13)
# 14: Soldes Débit (merged 14-15)
# 16: Soldes Crédit (merged 16-17)

# Print all rows where col 0 starts with 4, 5 or contains "TOTAL"
print("=" * 100)
print("BALANCE — Comptes 4xx et 5xx")
print("=" * 100)
print(f"{'N° compte':<10} {'Intitulé':<40} {'Mvt Débit':>18} {'Mvt Crédit':>18} {'Solde D':>15} {'Solde C':>15}")
print("-" * 116)

total_debit = 0.0
total_credit = 0.0
total_sdebit = 0.0
total_scredit = 0.0
nb_comptes = 0

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 12:  # skip header
        continue
    compte = str(row[0] or "").strip()
    if not compte:
        continue
    intitule = str(row[4] or "").strip()
    mvt_debit = float(row[9]) if row[9] is not None else 0.0
    mvt_credit = float(row[11]) if row[11] is not None else 0.0
    solde_debit = float(row[14]) if row[14] is not None else 0.0
    solde_credit = float(row[16]) if row[16] is not None else 0.0

    if compte.startswith("4") or compte.startswith("5") or "TOTAL" in compte.upper():
        print(f"{compte:<10} {intitule[:40]:<40} {mvt_debit:>18,.2f} {mvt_credit:>18,.2f} {solde_debit:>15,.2f} {solde_credit:>15,.2f}")

    # Look for totals
    if "TOTAL" in compte.upper() or "TOTAL" in intitule.upper():
        print(f"  >> TOTAUX FOUND at row {i+1}: D={mvt_debit:,.2f} C={mvt_credit:,.2f} SD={solde_debit:,.2f} SC={solde_credit:,.2f}")

    if compte[:3].isdigit():
        nb_comptes += 1
        total_debit += mvt_debit
        total_credit += mvt_credit
        total_sdebit += solde_debit
        total_scredit += solde_credit

print(f"\nComputed totals (sum of all accounts with numeric prefix):")
print(f"  Total Mvt Débit   = {total_debit:,.2f}")
print(f"  Total Mvt Crédit  = {total_credit:,.2f}")
print(f"  Équilibre mvt     = {total_debit - total_credit:,.2f} (doit être 0)")
print(f"  Total Solde Débit = {total_sdebit:,.2f}")
print(f"  Total Solde Crédit= {total_scredit:,.2f}")
print(f"  Équilibre soldes  = {total_sdebit - total_scredit:,.2f} (doit être 0)")
print(f"  Nombre de comptes = {nb_comptes}")

# Print last 20 rows to find totals row
print("\n--- Last 20 rows ---")
all_rows = list(ws.iter_rows(values_only=True))
for row in all_rows[-20:]:
    trimmed = [str(c)[:30] if c is not None else "" for c in row[:18]]
    if any(trimmed):
        idx = all_rows.index(row) + 1
        print(f"R{idx:>3}: {trimmed}")

wb.close()
