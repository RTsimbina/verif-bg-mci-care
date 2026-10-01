#!/usr/bin/env python3
"""Test /api/export endpoint with the JSON result (already analyzed)."""
import requests
import json
import os

# Load the previously-saved analysis result
with open("/home/z/my-project/scripts/analyze_result.json", "r") as f:
    result = json.load(f)

print(f"POST /api/export (JSON mode) ...")
r = requests.post(
    "http://localhost:3000/api/export",
    json=result,
    headers={"Content-Type": "application/json"},
    timeout=60,
)
print(f"Status: {r.status_code}")
print(f"Content-Type: {r.headers.get('Content-Type')}")
print(f"Content-Disposition: {r.headers.get('Content-Disposition')}")

if r.status_code != 200:
    print(r.text[:2000])
    raise SystemExit(1)

out_path = "/home/z/my-project/download/Export_Verif_BG_test.xlsx"
with open(out_path, "wb") as f:
    f.write(r.content)
print(f"Saved export to: {out_path} ({len(r.content):,} bytes)")

# Verify by re-reading
import openpyxl
wb = openpyxl.load_workbook(out_path, data_only=True)
print(f"\nSheets: {wb.sheetnames}")
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"  {sn}: {ws.max_row} rows x {ws.max_column} cols")
    # Print first 3 rows
    for i, row in enumerate(ws.iter_rows(max_row=3, values_only=True)):
        trimmed = []
        for c in row:
            s = str(c) if c is not None else ""
            if len(s) > 40:
                s = s[:37] + "..."
            trimmed.append(s)
        print(f"    R{i+1}: {trimmed}")
wb.close()
